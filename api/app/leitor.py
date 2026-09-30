"""Leitura estrutural das planilhas de ITBI-SP.

Responsabilidades (lado computador, seção 5 da especificação):
- abrir o XLSX em modo streaming, listar abas mensais, detectar cabeçalhos;
- associar colunas a nomes canônicos e registrar conflitos/colunas desconhecidas;
- converter cada célula para uma representação transportável em JSON, SEM limpar o valor.

O que NÃO é feito aqui: strip, CEP, conversão numérica, datas inválidas etc. — isso é
trabalho do ESP32. Strings seguem exatamente como vieram da célula.
"""
import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Callable, Iterator

import openpyxl

from .canonico import CABECALHOS, COLUNAS, aba_mensal, normalizar_cabecalho

LINHAS_BUSCA_CABECALHO = 10
MIN_COLUNAS_RECONHECIDAS = 3

# Marcadores de tipo da célula de origem, transportados no envelope MQTT (campo "tipos").
TIPO_NUMERO = "n"
TIPO_TEXTO = "s"
TIPO_DATA = "d"
TIPO_HORA = "t"
TIPO_BOOL = "b"


@dataclass
class ErroLeitura:
    aba: str
    linha: int | None
    tipo: str
    detalhe: dict


@dataclass
class MapaColunas:
    por_indice: dict[int, str]          # índice da coluna -> nome canônico
    extras: dict[int, str]              # índice -> cabeçalho original (desconhecido/conflito/vazio)
    linha_cabecalho: int | None         # número real (1-based) da linha de cabeçalho; None = sem cabeçalho
    erros: list[ErroLeitura] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


@dataclass
class RegistroLido:
    aba: str
    linha: int                          # número real da linha na planilha (1-based)
    dados: dict[str, str | None]        # 28 colunas canônicas
    tipos: dict[str, str]               # tipo da célula de origem, só para células não nulas
    extras: dict[str, str | None]       # colunas fora do esquema que tinham valor


def serializar_celula(valor) -> tuple[str | None, str | None]:
    """Converte uma célula openpyxl em (texto, tipo) sem alterar o significado.

    - números: texto decimal (sem notação científica; inteiros nunca passam por float);
    - datas: ISO 8601 (só a data quando o horário é 00:00:00);
    - strings: inalteradas (inclusive espaços), pois a limpeza é do ESP32.
    """
    if valor is None:
        return None, None
    if isinstance(valor, bool):
        return ("true" if valor else "false"), TIPO_BOOL
    if isinstance(valor, int):
        return str(valor), TIPO_NUMERO
    if isinstance(valor, float):
        texto = repr(valor)
        if "e" in texto or "E" in texto:
            texto = format(Decimal(texto), "f")
        return texto, TIPO_NUMERO
    if isinstance(valor, Decimal):
        return format(valor, "f"), TIPO_NUMERO
    if isinstance(valor, dt.datetime):
        if valor.time() == dt.time(0, 0):
            return valor.date().isoformat(), TIPO_DATA
        return valor.isoformat(), TIPO_DATA
    if isinstance(valor, dt.date):
        return valor.isoformat(), TIPO_DATA
    if isinstance(valor, dt.time):
        return valor.isoformat(), TIPO_HORA
    if isinstance(valor, str):
        return valor, TIPO_TEXTO
    return str(valor), TIPO_TEXTO


def listar_abas(caminho: str | Path) -> list[dict]:
    """Abas mensais do arquivo, com contagem aproximada de linhas (dimensão declarada)."""
    # Abrir via objeto de arquivo: o openpyxl recusa nomes sem extensão .xlsx (ex.: temporários).
    with open(caminho, "rb") as f:
        return _listar_abas(openpyxl.load_workbook(f, read_only=True, data_only=True))


def _listar_abas(wb) -> list[dict]:
    try:
        abas = []
        for ws in wb.worksheets:
            info = aba_mensal(ws.title)
            if not info:
                continue
            abas.append({
                "aba": ws.title,
                "mes": info[0],
                "ano": info[1],
                "linhas_aprox": max((ws.max_row or 1) - 1, 0),
            })
        return abas
    finally:
        wb.close()


def mapear_cabecalho(aba: str, celulas: tuple, linha: int | None) -> MapaColunas:
    nomes = [normalizar_cabecalho(c) for c in celulas]
    mapa = MapaColunas(por_indice={}, extras={}, linha_cabecalho=linha)

    # 2019: duas colunas "ACC (IPTU)". A primeira contém a descrição do padrão e a última o ACC.
    idx_acc = [i for i, n in enumerate(nomes) if CABECALHOS.get(n) == "acc_iptu"]
    tem_descricao = any(CABECALHOS.get(n) == "descricao_padrao" for n in nomes)
    resolvidos: dict[int, str] = {}
    if len(idx_acc) == 2 and not tem_descricao:
        resolvidos[idx_acc[0]] = "descricao_padrao"
        resolvidos[idx_acc[1]] = "acc_iptu"
        mapa.avisos.append(
            f"ACC (IPTU) duplicado: coluna {idx_acc[0] + 1} -> descricao_padrao, "
            f"coluna {idx_acc[1] + 1} -> acc_iptu"
        )

    usados: dict[str, int] = {}
    for i, nome in enumerate(nomes):
        canonico = resolvidos.get(i) or CABECALHOS.get(nome)
        if not nome:
            mapa.extras[i] = ""  # coluna sem cabeçalho; só vira erro se tiver dado
            continue
        if canonico is None:
            mapa.extras[i] = str(celulas[i])
            mapa.erros.append(ErroLeitura(aba, linha, "coluna_desconhecida",
                                          {"coluna": i + 1, "cabecalho": str(celulas[i])}))
            continue
        if canonico in usados:
            mapa.extras[i] = str(celulas[i])
            mapa.erros.append(ErroLeitura(aba, linha, "conflito_cabecalho", {
                "coluna": i + 1, "cabecalho": str(celulas[i]), "canonico": canonico,
                "coluna_mantida": usados[canonico] + 1,
            }))
            continue
        usados[canonico] = i
        mapa.por_indice[i] = canonico

    for c in COLUNAS:
        if c not in usados:
            mapa.erros.append(ErroLeitura(aba, linha, "coluna_ausente", {"canonico": c}))
    return mapa


def _detectar_cabecalho(aba: str, primeiras: list[tuple]) -> MapaColunas:
    for n, celulas in enumerate(primeiras, start=1):
        reconhecidas = sum(1 for c in celulas if normalizar_cabecalho(c) in CABECALHOS)
        if reconhecidas >= MIN_COLUNAS_RECONHECIDAS:
            return mapear_cabecalho(aba, celulas, n)
    # Sem cabeçalho: mesma estratégia posicional do código de referência (28 primeiras colunas).
    mapa = MapaColunas(por_indice={i: c for i, c in enumerate(COLUNAS)}, extras={}, linha_cabecalho=None)
    mapa.erros.append(ErroLeitura(aba, None, "aba_sem_cabecalho",
                                  {"estrategia": "posicional", "colunas": len(COLUNAS)}))
    return mapa


def ler_registros(
    caminho: str | Path,
    abas: list[str],
    limite: int | None = None,
    ao_erro: Callable[[ErroLeitura], None] | None = None,
    ao_mapa: Callable[[str, MapaColunas], None] | None = None,
) -> Iterator[RegistroLido]:
    """Percorre as abas pedidas, na ordem dada, produzindo registros de dados válidos.

    Linhas totalmente vazias são ignoradas (não contam para o limite). O número da linha
    é o número real na planilha, formando o identificador estável de origem.
    """
    ao_erro = ao_erro or (lambda e: None)
    arquivo = open(caminho, "rb")
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        emitidos = 0
        for nome_aba in abas:
            ws = wb[nome_aba]
            linhas = ws.iter_rows(values_only=True)
            primeiras = []
            for _ in range(LINHAS_BUSCA_CABECALHO):
                try:
                    primeiras.append(next(linhas))
                except StopIteration:
                    break
            mapa = _detectar_cabecalho(nome_aba, primeiras)
            for e in mapa.erros:
                ao_erro(e)
            if ao_mapa:
                ao_mapa(nome_aba, mapa)

            inicio = (mapa.linha_cabecalho or 0) + 1
            pendentes = [(n, c) for n, c in enumerate(primeiras, start=1) if n >= inicio]

            def todas():
                yield from pendentes
                yield from enumerate(linhas, start=len(primeiras) + 1)

            for numero, celulas in todas():
                if celulas is None or all(v is None for v in celulas):
                    continue
                dados = {c: None for c in COLUNAS}
                tipos: dict[str, str] = {}
                extras: dict[str, str | None] = {}
                for i, valor in enumerate(celulas):
                    texto, tipo = serializar_celula(valor)
                    canonico = mapa.por_indice.get(i)
                    if canonico:
                        dados[canonico] = texto
                        if tipo:
                            tipos[canonico] = tipo
                    elif texto is not None:
                        rotulo = mapa.extras.get(i) or f"coluna_{i + 1}"
                        extras[rotulo] = texto
                        ao_erro(ErroLeitura(nome_aba, numero, "valor_em_coluna_fora_do_esquema",
                                            {"coluna": i + 1, "cabecalho": rotulo, "valor": texto}))
                yield RegistroLido(nome_aba, numero, dados, tipos, extras)
                emitidos += 1
                if limite is not None and emitidos >= limite:
                    return
    finally:
        wb.close()
        arquivo.close()
